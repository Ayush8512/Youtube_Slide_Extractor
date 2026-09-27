import os
import re
import time
import uuid
import shutil
import zipfile
import subprocess
from fastapi import HTTPException
from PIL import Image, ImageOps, ImageChops, ImageStat, ImageEnhance, ImageDraw, ImageFont
import yt_dlp

from src.config import TEMP_DIR, HAS_PPTX, HAS_OCR, ensure_ffmpeg
from src.services.history_service import cleanup_file, save_history

if HAS_PPTX:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

if HAS_OCR:
    import pytesseract


def create_pdf_cover_page(width, height, title, author, subject, slide_count):
    canvas = Image.new('RGB', (width, height), (15, 23, 42))
    draw = ImageDraw.Draw(canvas)

    # Header bar
    draw.rectangle([0, 0, width, int(height * 0.15)], fill=(139, 92, 246))

    # Draw text
    title_text = f"STUDY NOTES: {title[:35]}"
    author_text = f"Author / Presenter: {author or 'YT Extractor Pro'}"
    subject_text = f"Subject: {subject or 'Video Lecture'}"
    count_text = f"Total Slides Extracted: {slide_count}"
    date_text = f"Generated On: {time.strftime('%Y-%m-%d')}"

    y = int(height * 0.3)
    draw.text((int(width * 0.08), y), title_text, fill=(255, 255, 255))
    draw.text((int(width * 0.08), y + 40), subject_text, fill=(168, 85, 247))
    draw.text((int(width * 0.08), y + 80), author_text, fill=(203, 213, 225))
    draw.text((int(width * 0.08), y + 120), count_text, fill=(59, 130, 246))
    draw.text((int(width * 0.08), y + 160), date_text, fill=(148, 163, 184))

    return canvas


def extract_slides_process(url: str, rate: str = "smart", sens: str = "0.15", start: str = None, end: str = None):
    if not shutil.which("ffmpeg"):
        ensure_ffmpeg()

    if not shutil.which("ffmpeg"):
        raise HTTPException(status_code=500, detail="FFmpeg is not installed on the system.")

    task_id = str(uuid.uuid4())[0:8]
    video_template = os.path.join(TEMP_DIR, f"{task_id}_vid.%(ext)s")

    dl_opts = {'format': 'bestvideo[height<=480]', 'outtmpl': video_template, 'noplaylist': True}

    if start or end:
        args = []
        if start:
            args.extend(['-ss', start])
        if end:
            args.extend(['-to', end])
        dl_opts['external_downloader_args'] = {'ffmpeg': args}

    try:
        with yt_dlp.YoutubeDL(dl_opts) as ydl:
            info = ydl.extract_info(url, download=True)
            video_file = ydl.prepare_filename(info)
            if not video_file.endswith('.webm') and not video_file.endswith('.mp4'):
                video_file = video_file.rsplit('.', 1)[0] + '.mp4'

        out_folder = os.path.join(TEMP_DIR, f"task_{task_id}")
        os.makedirs(out_folder, exist_ok=True)

        if rate == "smart":
            vf_filter = "fps=1"
            vsync = []
        else:
            vf_filter = f"fps=1/{rate}"
            vsync = []

        cmd = ["ffmpeg", "-y", "-i", video_file, "-filter:v", vf_filter] + vsync + [os.path.join(out_folder, "slide_%03d.jpg")]
        subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

        cleanup_file(video_file)
        files = sorted([f for f in os.listdir(out_folder) if f.endswith('.jpg')])

        if rate == "smart" and len(files) > 1:
            filtered_files = [files[0]]
            last_kept_path = os.path.join(out_folder, files[0])

            try:
                last_img = Image.open(last_kept_path).convert('L').resize((128, 128))
            except Exception:
                last_img = None

            threshold = float(sens) * 80

            for f in files[1:]:
                curr_path = os.path.join(out_folder, f)
                if not last_img:
                    filtered_files.append(f)
                    last_kept_path = curr_path
                    try:
                        last_img = Image.open(curr_path).convert('L').resize((128, 128))
                    except Exception:
                        pass
                    continue

                try:
                    curr_img = Image.open(curr_path).convert('L').resize((128, 128))
                    diff = ImageChops.difference(last_img, curr_img)
                    stat = ImageStat.Stat(diff)
                    mae = stat.mean[0]

                    if mae > threshold:
                        filtered_files.append(f)
                        last_img = curr_img
                    else:
                        os.remove(curr_path)
                except Exception:
                    filtered_files.append(f)
                    try:
                        last_img = Image.open(curr_path).convert('L').resize((128, 128))
                    except Exception:
                        pass

            files = filtered_files

        save_history({"title": info.get('title', 'Extracted Slides'), "date": time.strftime("%Y-%m-%d %H:%M"), "slides": len(files)})
        return {"task_id": task_id, "files": files}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


def generate_file_process(data):
    task_folder = os.path.join(TEMP_DIR, f"task_{data.task_id}")
    if not os.path.exists(task_folder) and data.format not in ["md", "txt"]:
        raise HTTPException(status_code=404, detail="Session expired")

    safe_title = re.sub(r'[\\\\/*?:"<>|]', "_", data.title)
    out_file = os.path.join(TEMP_DIR, f"{safe_title}_{data.task_id}.{data.format}")

    images = []
    if os.path.exists(task_folder):
        for f in data.selected_files:
            p = os.path.join(task_folder, f)
            if os.path.exists(p):
                img = Image.open(p).convert('RGB')
                if data.invert_colors:
                    img = ImageOps.invert(img)
                if data.enhance_text:
                    enhancer = ImageEnhance.Contrast(img)
                    img = enhancer.enhance(1.5)
                    sharpener = ImageEnhance.Sharpness(img)
                    img = sharpener.enhance(2.0)
                images.append(img)

    if data.format == "pdf":
        if not images:
            raise HTTPException(status_code=400, detail="No valid images for PDF")

        cover = create_pdf_cover_page(images[0].width, images[0].height, data.title, data.author_name, data.subject_tag, len(data.selected_files))
        pdf_images = [cover] + images

        if data.layout > 1:
            processed = [cover]
            for i in range(0, len(images), data.layout):
                batch = images[i:i+data.layout]
                bg_w, bg_h = images[0].size
                if data.layout == 2:
                    canvas = Image.new('RGB', (bg_w, bg_h * 2), (255, 255, 255))
                    canvas.paste(batch[0], (0, 0))
                    if len(batch) > 1:
                        canvas.paste(batch[1], (0, bg_h))
                else:
                    canvas = Image.new('RGB', (bg_w * 2, bg_h * 2), (255, 255, 255))
                    canvas.paste(batch[0], (0, 0))
                    if len(batch) > 1:
                        canvas.paste(batch[1], (bg_w, 0))
                    if len(batch) > 2:
                        canvas.paste(batch[2], (0, bg_h))
                    if len(batch) > 3:
                        canvas.paste(batch[3], (bg_w, bg_h))
                processed.append(canvas)
            pdf_images = processed

        pdf_images[0].save(out_file, save_all=True, append_images=pdf_images[1:])
        mt = 'application/pdf'

    elif data.format == "zip":
        if not images:
            raise HTTPException(status_code=400, detail="No valid images for ZIP")
        with zipfile.ZipFile(out_file, 'w') as zipf:
            for idx, img in enumerate(images):
                temp_img = os.path.join(TEMP_DIR, f"temp_{idx}.jpg")
                img.save(temp_img)
                zipf.write(temp_img, arcname=data.selected_files[idx])
                os.remove(temp_img)
        mt = 'application/zip'

    elif data.format == "txt":
        text_content = f"--- Extracted Notes: {data.title} ---\\n\\n"
        if HAS_OCR and images:
            for idx, img in enumerate(images):
                f_name = data.selected_files[idx]
                note = data.slide_notes.get(f_name, "").strip()
                tag = data.slide_tags.get(f_name, "").strip()
                text_content += f"--- Slide {idx+1} ({f_name}) ---\\n"
                if tag:
                    text_content += f"Tag: {tag}\\n"
                if note:
                    text_content += f"Note: {note}\\n\\n"
                try:
                    text = pytesseract.image_to_string(img)
                    text_content += text.strip() + "\\n\\n"
                except Exception as e:
                    text_content += f"[OCR Error: {str(e)}]\\n\\n"
        else:
            text_content += f"Slide Count: {len(data.selected_files)}\\n"

        with open(out_file, "w", encoding="utf-8") as f:
            f.write(text_content)
        mt = "text/plain"

    elif data.format == "md":
        md_content = f"# 📚 Study Notes: {data.title}\\n\\n"
        md_content += f"**Total Extracted Slides:** {len(data.selected_files)}\\n\\n"
        md_content += "## 📌 Slide Overview & Custom Notes\\n\\n"
        for idx, f in enumerate(data.selected_files):
            md_content += f"### Slide #{idx+1} (`{f}`)\\n"
            tag = data.slide_tags.get(f, "").strip()
            note = data.slide_notes.get(f, "").strip()
            if tag:
                md_content += f"🏷️ **Tag**: `{tag}`\\n\\n"
            if note:
                md_content += f"> ✍️ **Custom Note**: {note}\\n\\n"
            if HAS_OCR and idx < len(images):
                try:
                    ocr_txt = pytesseract.image_to_string(images[idx]).strip()
                    if ocr_txt:
                        md_content += f"```text\\n{ocr_txt}\\n```\\n\\n"
                except Exception:
                    pass

        with open(out_file, "w", encoding="utf-8") as f:
            f.write(md_content)
        mt = "text/markdown"

    elif data.format == "pptx":
        if not HAS_PPTX:
            raise HTTPException(status_code=400, detail="python-pptx not installed")
        if not images:
            raise HTTPException(status_code=400, detail="No valid images for PPTX")
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = data.title
        title_slide.placeholders[1].text = f"Extracted Slides Notes • {time.strftime('%Y-%m-%d')}"

        for img in images:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            temp_img = os.path.join(TEMP_DIR, f"pptx_temp_{uuid.uuid4().hex[:6]}.jpg")
            img.save(temp_img)
            slide.shapes.add_picture(temp_img, 0, 0, width=prs.slide_width, height=prs.slide_height)
            os.remove(temp_img)
        prs.save(out_file)
        mt = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
    else:
        raise HTTPException(status_code=400, detail="Unsupported format")

    return out_file, mt
